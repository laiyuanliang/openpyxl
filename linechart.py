from openpyxl import Workbook
from openpyxl.chart import LineChart,Reference
from openpyxl.chart.axis import DateAxis
from datetime import date
from copy import deepcopy

wb = Workbook()
ws = wb.active

rows = [
    ['Date', 'Batch 1', 'Batch 2', 'Batch 3'],
    [date(2015,9, 1), 40, 30, 25],
    [date(2015,9, 2), 40, 25, 30],
    [date(2015,9, 3), 50, 30, 45],
    [date(2015,9, 4), 30, 25, 40],
    [date(2015,9, 5), 25, 35, 30],
    [date(2015,9, 6), 20, 40, 35],
]

for row in rows:
    ws.append(row)

c1 = LineChart()
c1.title = 'Line Chart'
c1.x_axis.title = 'number'
c1.y_axis.title = 'size'

data = Reference(ws, min_col =2, min_row =1, max_col =4, max_row =7)
c1.add_data(data, titles_from_data =True)

#style the line
s0 = c1.series[0]
s0.marker.symbol = 'triangle'
s0.marker.graphicalProperties.solidFill = 'ff0000' #seting color of marker
s0.marker.graphicalProperties.line.solidFill = 'ff0000' 
s0.graphicalProperties.line.noFill = True #hiding connect line

s1 = c1.series[1]
s1.graphicalProperties.line.solidFill = '00AAAA'
s1.graphicalProperties.line.dashStyle = 'sysDot'
s1.graphicalProperties.line.width = '50050'

s2 = c1.series[2]
s2.smooth = True

ws.add_chart(c1,'A10')

stack = deepcopy(c1)
stack.grouping = 'stacked'
stack.title = 'stacked line chart'
ws.add_chart(stack, 'A18')

c2 = LineChart()
c2.title = 'Date Axis'
c2.y_axis.title = 'Size'
c2.x_axis.title = 'Date'
c2.x_axis.majorTimeUnit = 'days'
c2.x_axis.number_format = 'd-mmm'
#c2.y_axis.crossAx = 500  # if add this line or next line, Excel couldn't open output file
#c2.x_axis = DateAxis(crossAx=100)
c2.add_data(data, titles_from_data =True)
dates = Reference(ws, min_col =1, min_row =2, max_row =7)
c2.set_categories(dates)
ws.add_chart(c2, 'A35')

wb.save('linechart.xlsx')
