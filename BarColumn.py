from openpyxl import Workbook
from openpyxl.chart import BarChart,Reference
from copy import deepcopy

wb = Workbook()
ws = wb.active
ws.title = 'BarColumnChart'

rows = [['uid', 'age', 'score'],
        [23, 18, 89],
        [17, 16, 80],
        [11, 20, 88],
        [7, 22, 60],
        [29, 20, 99],
        [30, 18, 98]
]

for row in rows:
    ws.append(row)

chart1 = BarChart()
chart1.title = 'Bar and Column Chart'
chart1.style = 10
chart1.type = 'col'
chart1.x_axis.title = 'uid'
chart1.y_axis.title = 'age and score'

data1 = Reference(ws, min_col =2, min_row =1, max_col =3, max_row =7)
cats1 = Reference(ws, min_col =1, min_row =2, max_row =7)

chart1.add_data(data1, titles_from_data =True)
chart1.set_categories(cats1)

ws.add_chart(chart1, 'A10')

chart2 = deepcopy(chart1)
chart2.title = 'Horizontal chart'
chart2.style = 11
chart2.type = 'bar'
ws.add_chart(chart2, 'G10')

chart3 = deepcopy(chart1)
chart3.title = 'column stack'
chart3.style = 12
chart3.type = 'col'
chart3.overlap = 100
chart3.grouping = 'stacked'
ws.add_chart(chart3, 'A18')


chart4 = deepcopy(chart1)
chart4.title = 'bar stack percent'
chart4.style = 13
chart4.type = 'bar'
chart4.overlap = 100
chart4.grouping = 'percentStacked'
ws.add_chart(chart4, 'G18')

wb.save('BarColumn.xlsx')
